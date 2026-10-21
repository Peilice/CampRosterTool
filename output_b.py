"""類型 B：學員狀況總覽.docx"""
import os

import docx
from docx.shared import Emu, Inches, Pt, RGBColor

from processing import SESSION_LABEL, group_students, is_blank_note

CARDS_PER_PAGE = 4
COLS = 2
ROWS = CARDS_PER_PAGE // COLS

LABEL_SIZE = Pt(18)
VALUE_SIZE = Pt(20)
NOTE_SIZE = Pt(16)
NOTE_COLOR = RGBColor(0xFF, 0x00, 0x00)

PAGE_WIDTH = Inches(8.27)
MARGIN = Inches(0.5)
USABLE_WIDTH = PAGE_WIDTH - 2 * MARGIN
COL_WIDTH = Emu(int(USABLE_WIDTH / COLS))


def _add_labeled_run(p, label, value="", value_size=None, value_bold=False, value_color=None):
    run = p.add_run(label)
    run.font.size = LABEL_SIZE
    if value:
        vrun = p.add_run(value)
        vrun.font.size = value_size or LABEL_SIZE
        vrun.font.bold = value_bold
        if value_color:
            vrun.font.color.rgb = value_color


def _write_card(cell, s):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    _add_labeled_run(p, "編號：", s.code, value_size=VALUE_SIZE, value_bold=True)

    p = cell.add_paragraph()
    _add_labeled_run(p, "姓名：", s.get("學員姓名"), value_size=VALUE_SIZE, value_bold=True)

    p = cell.add_paragraph()
    _add_labeled_run(p, "隊名：")

    note = "" if is_blank_note(s.get("學員特殊狀況說明")) else s.get("學員特殊狀況說明")
    p = cell.add_paragraph()
    _add_labeled_run(p, "注意事項：", note, value_size=NOTE_SIZE, value_color=NOTE_COLOR)

    cell.add_paragraph()
    cell.add_paragraph()

    p = cell.add_paragraph()
    _add_labeled_run(p, "備註：")


def _set_narrow_margins(section):
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN


def _write_group(doc, students):
    if not students:
        return
    for chunk_start in range(0, len(students), CARDS_PER_PAGE):
        chunk = students[chunk_start:chunk_start + CARDS_PER_PAGE]
        table = doc.add_table(rows=ROWS, cols=COLS)
        table.style = "Table Grid"
        table.autofit = False
        for row in table.rows:
            for cell in row.cells:
                cell.width = COL_WIDTH
        for i, s in enumerate(chunk):
            cell = table.rows[i // COLS].cells[i % COLS]
            cell.width = COL_WIDTH
            _write_card(cell, s)
        if chunk_start + CARDS_PER_PAGE < len(students):
            doc.add_page_break()


def build(students_all, sessions_wanted, folder, base_name):
    """學員狀況總覽：依梯次、性別各自產生獨立檔案；沒有學員的組別不產生。"""
    students = students_all["students"]
    groups = group_students(students, sessions_wanted)

    output_paths = []
    for session in sorted(sessions_wanted):
        label = SESSION_LABEL[session]
        for gender in ("男", "女"):
            group = groups.get((session, gender), [])
            if not group:
                continue

            doc = docx.Document()
            _set_narrow_margins(doc.sections[0])
            _write_group(doc, group)

            output_path = os.path.join(folder, f"{base_name}_第{label}梯_{gender}.docx")
            doc.save(output_path)
            output_paths.append(output_path)

    return output_paths
