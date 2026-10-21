"""類型 A：學員整理.xlsx"""
import os

import openpyxl
from openpyxl.styles import Border, PatternFill, Side

from processing import (
    SESSION_LABEL, group_students, mask_id, mask_name,
)
from xlsx_style import CONFLICT_FILL, style_header_row

DROP_COLUMNS = {"縣市", "郵遞區號"}

LIST_HEADERS = ["編號", "項次", "學員姓名", "出生年月日", "身分證字號", "管理者註記事項"]

BLOCK_WIDTH = 4  # 編號 / 姓名 / 身分證字號 / 空白

GRID_FILL_COLOR = {
    (1, "男"): "CCFFCC",  # 淺綠
    (1, "女"): "FFFFCC",  # 淺黃
    (2, "男"): "CCECFF",  # 淺藍
    (2, "女"): "FFCCCC",  # 淺紅
}

GRID_BORDER = Border(
    left=Side(style="thin", color="FF000000"),
    right=Side(style="thin", color="FF000000"),
    top=Side(style="thin", color="FF000000"),
    bottom=Side(style="thin", color="FF000000"),
)


def _write_raw_sheet(wb, header, rows, conflict_rows):
    ws = wb.active
    ws.title = "原檔"
    kept_headers = [h for h in header if h not in DROP_COLUMNS]
    ws.append(kept_headers)
    style_header_row(ws, 1, len(kept_headers))
    for i, row in enumerate(rows):
        raw = dict(zip(header, row))
        row_i = i + 2
        ws.append([raw.get(h, "") for h in kept_headers])
        if i in conflict_rows:
            for col in range(1, len(kept_headers) + 1):
                ws.cell(row=row_i, column=col).fill = CONFLICT_FILL


def _write_list_sheet(wb, sheet_name, students):
    ws = wb.create_sheet(sheet_name)
    ws.append(LIST_HEADERS)
    style_header_row(ws, 1, len(LIST_HEADERS))
    for s in students:
        ws.append([
            s.code, s.get("項次"), s.get("學員姓名"), s.birth_ymd,
            s.get("身分證字號"), s.get("管理者註記事項"),
        ])


def _write_grid_sheet(wb, sheet_name, title, students, masked, fill_color=None):
    ws = wb.create_sheet(sheet_name)
    total = len(students)
    block_count = max(1, -(-total // 50)) if total else 1
    ws.cell(row=1, column=5, value=title)
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid") if fill_color else None
    for i, s in enumerate(students):
        block = i // 50
        pos_in_block = i % 50
        row = 2 + pos_in_block
        col = block * BLOCK_WIDTH + 1
        name = mask_name(s.get("學員姓名")) if masked else s.get("學員姓名")
        id_no = mask_id(s.get("身分證字號")) if masked else s.get("身分證字號")
        for c, value in ((col, s.code), (col + 1, name), (col + 2, id_no)):
            cell = ws.cell(row=row, column=c, value=value)
            cell.border = GRID_BORDER
            if fill:
                cell.fill = fill

    return block_count


def build(students_all, sessions_wanted, folder, base_name):
    """學員整理：只產生一個檔案，梯次、性別各自分開分頁。"""
    header = students_all["header"]
    rows = students_all["rows"]
    students = students_all["students"]
    conflict_rows = students_all["conflict_rows"]

    wb = openpyxl.Workbook()
    _write_raw_sheet(wb, header, rows, conflict_rows)

    groups = group_students(students, sessions_wanted)
    for session in sorted(sessions_wanted):
        label = SESSION_LABEL[session]
        for gender, gender_tag in (("男", "男"), ("女", "女")):
            group = groups.get((session, gender), [])
            list_name = f"第{label}梯-{gender_tag}"
            grid_name = f"第{label}({gender_tag})"
            public_name = f"{label}公({gender_tag})"
            title = f"第{label}梯次({gender_tag})"
            fill_color = GRID_FILL_COLOR[(session, gender)]

            _write_list_sheet(wb, list_name, group)
            _write_grid_sheet(wb, grid_name, title, group, masked=False, fill_color=fill_color)
            _write_grid_sheet(wb, public_name, title, group, masked=True, fill_color=fill_color)

    output_path = os.path.join(folder, f"{base_name}.xlsx")
    wb.save(output_path)
    return [output_path]
