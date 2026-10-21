"""類型 C：報到及查詢資料.xlsx"""
import os

import openpyxl

from processing import SESSION_LABEL, group_students, roc_year_now
from xlsx_style import style_header_row

QUERY_HEADERS = ["編號", "姓名", "出生", "月", "日", "身分證字號", "地址", "就讀學校", "就讀年級", "緊急連絡人", "緊急連絡人電話"]
CHECKIN_HEADERS = ["報到", "編號", "姓名", "年次", "月", "日", "身分證字號", "住址"]


def _write_query_sheet(wb, sheet_name, title, students):
    ws = wb.create_sheet(sheet_name)
    ws.append([title])
    ws.append(QUERY_HEADERS)
    style_header_row(ws, 2, len(QUERY_HEADERS))
    for s in students:
        ws.append([
            s.code, s.get("學員姓名"), s.get("出生年"), s.get("月"), s.get("日"),
            s.get("身分證字號"), s.get("地址"), s.get("學員目前就讀學校"),
            s.get("學員就讀年級"), s.get("緊急連絡人姓名"), s.get("緊急連絡人電話"),
        ])


def _write_checkin_sheet(wb, sheet_name, title, students):
    ws = wb.create_sheet(sheet_name)
    ws.append([title])
    ws.append(CHECKIN_HEADERS)
    style_header_row(ws, 2, len(CHECKIN_HEADERS))
    for s in students:
        ws.append([
            None, s.code, s.get("學員姓名"), s.get("出生年"), s.get("月"), s.get("日"),
            s.get("身分證字號"), s.get("地址"),
        ])


def build(students_all, sessions_wanted, folder, base_name):
    """報到及查詢資料：依梯次各自產生獨立檔案，性別分頁保留在同一檔案內。"""
    students = students_all["students"]
    groups = group_students(students, sessions_wanted)
    year = roc_year_now()

    output_paths = []
    for session in sorted(sessions_wanted):
        label = SESSION_LABEL[session]
        checkin_code = f"{year}{session:02d}"

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for gender in ("男", "女"):
            group = groups.get((session, gender), [])

            query_name = f"第{label}梯{gender}查詢"
            query_title = f"元亨寺{year}年暑期學佛營第{label}梯學員({gender})資料"
            _write_query_sheet(wb, query_name, query_title, group)

            checkin_name = f"{checkin_code}{gender}"
            checkin_title = f"元亨寺{year}年暑期學佛夏令營學員報到單(第{label}梯次{gender}生)"
            _write_checkin_sheet(wb, checkin_name, checkin_title, group)

        output_path = os.path.join(folder, f"{base_name}_第{label}梯.xlsx")
        wb.save(output_path)
        output_paths.append(output_path)

    return output_paths
