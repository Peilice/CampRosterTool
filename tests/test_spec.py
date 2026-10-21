"""依 docs/功能規格.md 的範例驗證處理規則與報表。所有資料皆為虛構。

執行：python -m unittest discover tests
"""
import json
import os
import sys
import tempfile
import unittest

import docx
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import output_a  # noqa: E402
import output_d  # noqa: E402
import processing as P  # noqa: E402
import web_entry  # noqa: E402
from reader import read_source  # noqa: E402

# §3.4 兩種固定格式（只取測試需要的欄位）
WEB_HEADER = [
    "項次", "學員姓名", "性別", "出生年", "月", "日", "身分證字號",
    "參加梯次", "學員特殊狀況說明", "管理者註記事項",
]
ONSITE_HEADER = [
    "時間戳記", "項次", "1.學員姓名", "2.性別", "3.生日：(民國年)/月/日（例：104/05/06）",
    "4. 身份證字號", "8.報名梯次", "\n學員特殊狀況說明",
]

S1 = "第一梯次：暑假後升國小五年級-國中二年級"
S1_WAIT = S1 + "(候補)"
S2 = "第二梯次(國曆7月22日~26日)"


def web(item="1", name="王小明", gender="男", birth=("104", "5", "6"), id_no=None,
        session=S1, special="", note=""):
    return [item, name, gender, *birth, id_no or f"A1{int(item or 0):08d}", session, special, note]


def onsite(item="000001", name="林現場", gender="男", birth="104/05/06", id_no=None,
           session=S1, special=""):
    return ["2026/7/1 上午 9:00:00", item, name, gender, birth,
            id_no or f"B1{abs(hash(item)) % 10**8:08d}", session, special]


def write_xlsx(folder, web_rows=(), onsite_rows=(), filename="source.xlsx", web_header=WEB_HEADER):
    """網路報名區塊從第 1 列開始，現場報名區塊接在後面（與實際原始檔相同）。

    沒有網路報名資料時，第 1 列仍是網路報名表頭，現場報名資料從第 3 列開始。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(web_header)
    for r in web_rows:
        ws.append(r)
    if onsite_rows:
        ws.append(ONSITE_HEADER)
        for r in onsite_rows:
            ws.append(r)
    path = os.path.join(folder, filename)
    wb.save(path)
    return path


class SpecTestCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = self._tmp.name

    def tearDown(self):
        self._tmp.cleanup()

    def process(self, web_rows=(), onsite_rows=()):
        header, data, row_nos = read_source(write_xlsx(self.tmp, web_rows, onsite_rows))
        result = P.build_students(header, data, row_nos)
        if not any(result["blockers"].values()):
            P.assign_codes(result)
        result["header"], result["rows"] = header, data
        return result

    def by_name(self, result):
        return {s.get("學員姓名"): s for s in result["students"]}


class SourceTest(SpecTestCase):
    """§3.4 依表頭格式區分來源"""

    def test_blocks_by_header_format(self):
        result = self.process([web(name="網路生")], [onsite(name="現場生", item="000003")])
        s = self.by_name(result)
        self.assertEqual(s["網路生"].source_type, "網路")
        self.assertEqual((s["網路生"].row_no, s["網路生"].code), (2, "000051"))
        self.assertEqual(s["現場生"].source_type, "現場")
        self.assertEqual((s["現場生"].row_no, s["現場生"].code), (4, "000003"))

    def test_onsite_aliases_mapped(self):
        result = self.process(onsite_rows=[onsite(special="對花生過敏", session="第二梯次", gender="女",
                                                  item="300001")])
        s = result["students"][0]
        self.assertEqual((s.session, s.gender), (2, "女"))
        self.assertEqual(s.get("學員特殊狀況說明"), "對花生過敏")
        self.assertTrue(s.get("身分證字號"))


class BirthTest(SpecTestCase):
    """§3.5 出生年月日"""

    def test_examples(self):
        cases = [
            (("104", "5", "6"), "104/05/06"),
            (("104/05/06", "", ""), "104/05/06"),
            (("104-5-6", "", ""), "104/05/06"),
            (("104／05／06", "", ""), "104/05/06"),
            (("1040506", "", ""), "104/05/06"),
            (("2015/05/06", "", ""), "104/05/06"),
            (("2015", "5", "6"), "104/05/06"),
        ]
        rows = [web(item=str(i), name=f"學員{i}", birth=b) for i, (b, _) in enumerate(cases, start=1)]
        result = self.process(rows)
        students = self.by_name(result)
        for i, (_, expected) in enumerate(cases, start=1):
            self.assertEqual(students[f"學員{i}"].birth_ymd, expected, cases[i - 1])
        self.assertEqual(result["issues"]["bad_birth"], [])

    def test_onsite_combined_birth(self):
        result = self.process(onsite_rows=[onsite(birth="104/5/6")])
        self.assertEqual(result["students"][0].birth_ymd, "104/05/06")

    def test_unparseable_keeps_value_and_warns(self):
        result = self.process([web(birth=("104年5月", "", ""))])
        self.assertEqual(result["students"][0].get("出生年"), "104年5月")
        self.assertEqual(result["issues"]["bad_birth"], ["第 2 列 王小明：出生年月日「104年5月」"])


class ExclusionTest(SpecTestCase):
    """§4.1 排除"""

    def test_examples(self):
        result = self.process(
            [web(item="1", name="甲", session="第一梯次(國曆7月7日~11日)"),
             web(item="2", name="丙", gender="南", session="第二梯次")],
            [onsite(name="乙", gender="女", session="")],
        )
        self.assertEqual(list(self.by_name(result)), ["甲"])
        self.assertEqual(result["unrecognized_rows"], {1, 2})
        self.assertEqual(result["issues"]["unrecognized"], [
            "第 3 列 丙：性別「南」無法辨識",
            "第 5 列 乙：參加梯次空白",
        ])

    def test_leading_zeros_restored(self):
        # 試算表把 000012 存成 12：第一梯男補回後相符，其他組別補回後組別碼不符
        result = self.process(onsite_rows=[
            onsite(item="12", name="第一梯男"),
            onsite(item="13", name="第一梯女", gender="女"),
        ])
        self.assertEqual(self.by_name(result)["第一梯男"].code, "000012")
        self.assertEqual(result["issues"]["conflict"], ["第 4 列 第一梯女：編號 000013 與第一梯女不符"])
        self.assertEqual(result["blockers"]["invalid_onsite_no"], [])

    def test_onsite_code_conflict(self):
        result = self.process(onsite_rows=[
            onsite(item="200012", name="丁", gender="女", session="第二梯次"),
            onsite(item="300013", name="戊", gender="女", session="第二梯次"),
        ])
        self.assertEqual(result["conflict_rows"], {0})
        self.assertEqual(result["issues"]["conflict"], ["第 3 列 丁：編號 200012 與第二梯女不符"])
        self.assertEqual(self.by_name(result)["戊"].code, "300013")


class NumberingTest(SpecTestCase):
    """§4.2 狀態、§5 編號"""

    def test_spec_example_first_session_female(self):
        f = dict(gender="女")
        result = self.process(
            [
                web(name="網路158", item="158", **f),
                web(name="網路203", item="203", note="取消 家長來電", **f),
                web(name="網路640", item="640", **f),
                web(name="網路655", item="655", session=S1_WAIT, note="轉正2", **f),
                web(name="網路700", item="700", session=S1_WAIT, note="轉正1", **f),
                web(name="網路710", item="710", session=S1_WAIT, **f),
            ],
            [
                onsite(name="現場12", item="100012", **f),
                onsite(name="現場3", item="100003", **f),
            ],
        )
        s = self.by_name(result)
        expected = {
            "現場3": ("100003", P.ADMITTED),
            "現場12": ("100012", P.ADMITTED),
            "網路158": ("100051", P.ADMITTED),
            "網路203": ("100052", P.CANCELLED),
            "網路640": ("100053", P.ADMITTED),
            "網路700": ("100054", P.PROMOTED),
            "網路655": ("100055", P.PROMOTED),
            "網路710": (None, P.WAITLISTED),
        }
        for name, (code, status) in expected.items():
            self.assertEqual((s[name].code, s[name].status), (code, status), name)
        # §5.3 排序：現場在前、網路在後，候補排最後
        self.assertEqual([x.get("學員姓名") for x in result["students"]], [
            "現場3", "現場12", "網路158", "網路203", "網路640", "網路700", "網路655", "網路710",
        ])

    def test_cancel_after_promotion_keeps_number(self):
        result = self.process([
            web(name="正取", item="1"),
            web(name="先轉正後取消", item="2", session=S1_WAIT, note="取消 轉正1"),
            web(name="後轉正", item="3", session=S1_WAIT, note="轉正2"),
        ])
        s = self.by_name(result)
        self.assertEqual(s["先轉正後取消"].code, "000052")
        self.assertEqual(s["先轉正後取消"].status, P.CANCELLED)
        self.assertEqual(s["後轉正"].code, "000053")

    def test_marker_only_at_start(self):
        result = self.process([web(note="午餐取消")])
        self.assertEqual(result["students"][0].status, P.ADMITTED)

    def test_promote_marker_without_waitlist_warns(self):
        result = self.process([web(note="轉正1")])
        self.assertEqual(result["students"][0].status, P.ADMITTED)
        self.assertEqual(len(result["issues"]["promote_not_waitlisted"]), 1)

    def test_onsite_over_50_warns(self):
        result = self.process(onsite_rows=[onsite(item="100053", gender="女")])
        self.assertEqual(result["students"][0].code, "100053")
        self.assertEqual(len(result["issues"]["onsite_over_50"]), 1)

    def test_duplicate_id_warns(self):
        result = self.process([
            web(name="陳小美", item="1", id_no="a100000001"),
            web(name="陳小美", item="2", id_no="A100000001 "),
        ])
        self.assertEqual(result["issues"]["duplicate_id"], ["第 2 列 陳小美、第 3 列 陳小美：身分證字號相同"])


class BlockerTest(SpecTestCase):
    """§4.3、§7.3 不准產生"""

    def test_invalid_onsite_code(self):
        result = self.process(onsite_rows=[
            onsite(item="", name="甲"),
            onsite(item="1000012", name="乙"),
            onsite(item="12號", name="丙"),
        ])
        self.assertEqual(result["blockers"]["invalid_onsite_no"], [
            "第 3 列 甲：項次空白",
            "第 4 列 乙：項次「1000012」不是 6 碼完整編號",
            "第 5 列 丙：項次「12號」不是 6 碼完整編號",
        ])

    def test_duplicate_onsite_code(self):
        result = self.process(onsite_rows=[
            onsite(item="000005", name="甲"),
            onsite(item="000005", name="乙"),
            onsite(item="100005", name="丙", gender="女"),
        ])
        self.assertEqual(result["blockers"]["duplicate_onsite_no"], ["第 3 列 甲、第 4 列 乙：編號都是 000005"])

    def test_duplicate_promote_n(self):
        result = self.process([
            web(name="甲", item="1", session=S1_WAIT, note="轉正1"),
            web(name="乙", item="2", session=S1_WAIT, note="轉正1"),
        ])
        self.assertEqual(result["blockers"]["duplicate_promote_n"], ["第 2 列 甲、第 3 列 乙：都標記為轉正1"])


class MaskTest(unittest.TestCase):
    """§6.1 遮蔽規則"""

    def test_examples(self):
        self.assertEqual(P.mask_name("林"), "林")
        self.assertEqual(P.mask_name("王明"), "王○")
        self.assertEqual(P.mask_name("陳小美"), "陳○美")
        self.assertEqual(P.mask_name("歐陽小明"), "歐○○明")
        self.assertEqual(P.mask_id("A123456789"), "A12****789")


class ReportTest(SpecTestCase):
    """§6 報表"""

    def test_member_list_shows_all_statuses_and_grids_only_reportable(self):
        data = self.process([
            web(name="正取生", item="1"),
            web(name="候補生", item="2", session=S1_WAIT),
            web(name="取消生", item="3", note="取消"),
            web(name="未辨識", item="4", session=""),
        ])
        path = output_a.build(data, {1}, self.tmp, "學員整理")[0]
        wb = openpyxl.load_workbook(path)
        listed = [(r[2], r[6]) for r in wb["第一梯-男"].iter_rows(min_row=2, values_only=True)]
        self.assertEqual(listed, [("正取生", None), ("取消生", "取消"), ("候補生", "候補")])
        public_names = [c.value for c in wb["一公(男)"]["B"] if c.value]
        self.assertEqual(public_names, ["正○生"])
        raw = wb["原檔"]
        self.assertEqual(raw.cell(row=5, column=1).fill.fgColor.rgb, "FFFFC000")  # 未辨識標橘色
        self.assertEqual(raw.cell(row=2, column=1).value, "網路報名")  # 來源依格式填入

    def test_formation_notes_combine_both_fields(self):
        data = self.process([
            web(name="甲", item="1", special="對花生過敏", note="家長要求和哥哥同隊"),
            web(name="乙", item="2", special="無", note="--"),
        ])
        path = output_d.build(data, {1}, self.tmp, "編隊")[0]
        wb = openpyxl.load_workbook(path)
        raw = list(wb["第一梯-男-原始資料"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(raw[0][4:], ("家長要求和哥哥同隊", "對花生過敏"))
        self.assertEqual(raw[1][4:], (None, None))  # 空白值不顯示
        formula = wb["第一梯-男-編隊"]["F2"].value
        self.assertIn("；", formula)
        self.assertIn("A:F,6,FALSE", formula)


class WebEntryTest(SpecTestCase):
    """§7 網頁操作"""

    def setUp(self):
        super().setUp()
        web_entry.OUT_DIR = os.path.join(self.tmp, "out")

    def load(self, web_rows=(), onsite_rows=(), **kwargs):
        return json.loads(web_entry.load(write_xlsx(self.tmp, web_rows, onsite_rows, **kwargs)))

    def test_blocker_disables_generation(self):
        info = self.load(onsite_rows=[onsite(item="", name="王小明")])
        self.assertFalse(info["ok"])
        self.assertEqual(info["messages"][0]["kind"], "error")
        self.assertEqual(info["messages"][0]["items"], ["第 3 列 王小明：項次空白"])

    def test_no_header_is_friendly_error(self):
        info = self.load(web_header=["姓名", "電話"])
        self.assertFalse(info["ok"])
        self.assertIn("找不到報名資料", info["messages"][0]["text"])

    def test_summary_and_empty_docx_group_skipped(self):
        info = self.load([
            web(name="甲", item="1"),
            web(name="乙", item="2", session=S1_WAIT),
        ])
        self.assertTrue(info["ok"])
        self.assertEqual((info["total"], info["included"]), (2, 1))
        self.assertEqual(info["status_counts"], {"候補": 1})
        out = json.loads(web_entry.generate("2", "1", "總覽"))
        self.assertEqual([os.path.basename(p) for p in out["files"]], ["總覽_第一梯_男.docx"])
        self.assertEqual(out["notices"], ["第一梯女沒有學員，未產生檔案。"])
        self.assertEqual(len(docx.Document(out["files"][0]).tables), 1)

    def test_file_name_cleanup(self):
        self.load([web()])
        out = json.loads(web_entry.generate("1", "1", "113/學員:整理.xlsx"))
        self.assertEqual(os.path.basename(out["files"][0]), "113_學員_整理.xlsx")
        out = json.loads(web_entry.generate("1", "1", " / "))
        self.assertEqual(os.path.basename(out["files"][0]), "學員整理.xlsx")


class SpreadsheetMLTest(SpecTestCase):
    """§3.1 網路報名系統匯出的 XML 格式 .xls，含 ss:Index 跳列"""

    def test_row_numbers_follow_index(self):
        cells = "".join(f'<Cell><Data ss:Type="String">{h}</Data></Cell>' for h in WEB_HEADER)
        values = "".join(f'<Cell><Data ss:Type="String">{v}</Data></Cell>' for v in web())
        xml = (
            '<?xml version="1.0"?>'
            '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
            'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
            '<Worksheet ss:Name="Registration"><Table>'
            f'<Row>{cells}</Row><Row ss:Index="5">{values}</Row>'
            '</Table></Worksheet></Workbook>'
        )
        path = os.path.join(self.tmp, "source.xls")
        with open(path, "w", encoding="utf-8") as f:
            f.write(xml)
        _, _, row_nos = read_source(path)
        self.assertEqual(row_nos, [5])


if __name__ == "__main__":
    unittest.main()
