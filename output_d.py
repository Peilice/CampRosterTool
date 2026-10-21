"""類型 D：編隊.xlsx"""
import os

import openpyxl
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule

from processing import SESSION_LABEL, group_students
from xlsx_style import style_header_row

SLOTS_PER_TEAM = 8

TEAM_NAMES = {
    "男": [
        "大願隊", "精進隊", "金剛隊", "莊嚴隊", "正覺隊", "大雄隊", "神通隊", "無畏隊",
        "法輪隊", "讚歎隊", "菩提隊", "天台隊", "龍王隊", "發心隊", "威德隊", "吉祥隊",
        "慧日隊", "堅固隊", "功德隊", "無量隊", "法界隊", "慧炬隊", "虛空隊", "不動隊",
        "歡喜隊", "大智隊", "海德隊", "聖賢隊", "安樂隊", "梵天隊", "聲聞隊", "緣覺隊",
        "大士隊", "華嚴隊", "無盡隊", "寶藏隊",
    ],
    "女": [
        "法音隊", "清淨隊", "真如隊", "大千隊", "普陀隊", "蓮華隊", "大悲隊", "微妙隊",
        "喜捨隊", "琉璃隊", "解脫隊", "甘露隊", "光明隊", "法雨隊", "究竟隊", "圓滿隊",
        "自在隊", "因緣隊", "摩尼隊", "普光隊", "善意隊", "金華隊", "妙音隊", "日光隊",
        "月光隊", "寶華隊", "慈音隊", "法雲隊", "妙樂隊", "滿意隊", "寶月隊", "寶燈隊",
    ],
}

RAW_HEADERS = ["編號", "學員姓名", "出生年月日", "身分證字號", "管理者註記事項"]

INPUT_FILL = PatternFill("solid", fgColor="FFCFE2F3")
DUPLICATE_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
TEAM_BORDER = Border(bottom=Side(style="medium"))

INSTRUCTIONS = [
    "填入學員編號，即可帶入姓名、注意事項",
    "學員編號為6碼數字，第一碼代表梯次與性別（0=第一梯男、1=第一梯女、2=第二梯男、3=第二梯女）",
    "表格有寫公式，若需增加每隊人數，插入後複製上面公式即可",
    "有底色代表，需要自己填入",
    "當日報名學員，在原始資料新增編號、姓名，並備註",
    "同一個學員編號填入超過一次，該儲存格會自動標示紅色，請人工核對",
]


def _write_team_name_sheet(wb, sheet_name, gender):
    ws = wb.create_sheet(sheet_name)
    ws.append(["性別", "序號", "隊名"])
    style_header_row(ws, 1, 3)
    for i, team_name in enumerate(TEAM_NAMES[gender], start=1):
        ws.append([gender, i, team_name])


def _write_raw_sheet(wb, sheet_name, students):
    ws = wb.create_sheet(sheet_name)
    ws.append(RAW_HEADERS)
    style_header_row(ws, 1, len(RAW_HEADERS))
    for s in students:
        ws.append([
            s.code, s.get("學員姓名"), s.birth_ymd,
            s.get("身分證字號"), s.get("管理者註記事項"),
        ])


def _write_formation_sheet(wb, sheet_name, gender, team_name_sheet, raw_sheet_name):
    ws = wb.create_sheet(sheet_name)
    ws.append(["隊名", "隊號", "隊員編號", "學員編號", "姓名", "注意事項"])
    style_header_row(ws, 1, 6)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 46
    ws.column_dimensions["H"].width = 6

    team_names = TEAM_NAMES[gender]
    row = 2
    for team_no in range(1, len(team_names) + 1):
        for slot in range(1, SLOTS_PER_TEAM + 1):
            ws.cell(row=row, column=1,
                    value=f"=VLOOKUP(B{row},{team_name_sheet}!B:C,2,FALSE)")
            ws.cell(row=row, column=2, value=team_no)
            ws.cell(row=row, column=3, value=slot)
            d_cell = ws.cell(row=row, column=4)
            d_cell.fill = INPUT_FILL
            d_cell.number_format = "@"
            ws.cell(row=row, column=5,
                    value=f"=IFERROR(VLOOKUP(D{row},'{raw_sheet_name}'!A:E,2,FALSE),\"\")")
            note_lookup = f"VLOOKUP(D{row},'{raw_sheet_name}'!A:E,5,FALSE)"
            ws.cell(row=row, column=6,
                    value=f"=IFERROR(IF({note_lookup}=0,\"\",{note_lookup}),\"\")")
            row += 1
        end_row = row - 1

        for col in range(1, 7):
            cell = ws.cell(row=end_row, column=col)
            cell.border = Border(
                left=cell.border.left, right=cell.border.right,
                top=cell.border.top, bottom=Side(style="medium"),
            )

    last_row = row - 1
    ws.conditional_formatting.add(
        f"D2:D{last_row}",
        FormulaRule(
            formula=[f'AND(D2<>"",COUNTIF($D$2:$D${last_row},D2)>1)'],
            fill=DUPLICATE_FILL,
        ),
    )

    ws.cell(row=2, column=8, value="說明")
    for i, text in enumerate(INSTRUCTIONS, start=3):
        ws.cell(row=i, column=8, value=i - 2)
        ws.cell(row=i, column=9, value=text)


def build(students_all, sessions_wanted, folder, base_name):
    """編隊：依梯次各自產生獨立檔案，性別分頁保留在同一檔案內。"""
    students = students_all["students"]
    groups = group_students(students, sessions_wanted)

    output_paths = []
    for session in sorted(sessions_wanted):
        label = SESSION_LABEL[session]

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        team_name_sheets = {}
        for gender in ("男", "女"):
            sheet_name = f"{gender}隊名"
            _write_team_name_sheet(wb, sheet_name, gender)
            team_name_sheets[gender] = sheet_name

        for gender in ("男", "女"):
            group = groups.get((session, gender), [])
            raw_name = f"第{label}梯-{gender}-原始資料"
            formation_name = f"第{label}梯-{gender}-編隊"
            _write_raw_sheet(wb, raw_name, group)
            _write_formation_sheet(
                wb, formation_name, gender, team_name_sheets[gender], raw_name,
            )

        output_path = os.path.join(folder, f"{base_name}_第{label}梯.xlsx")
        wb.save(output_path)
        output_paths.append(output_path)

    return output_paths
