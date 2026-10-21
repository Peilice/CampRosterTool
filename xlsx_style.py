"""共用的 Excel 表頭樣式設定。"""
from openpyxl.styles import Font, PatternFill

HEADER_FILL = PatternFill("solid", fgColor="FF1F3864")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)

# 編號與實際梯次/性別欄位衝突時的警示底色，提示人工核對。
CONFLICT_FILL = PatternFill("solid", fgColor="FFFFFF00")

# 梯次、性別或來源無法辨識而未納入報表的資料列底色。
UNRECOGNIZED_FILL = PatternFill("solid", fgColor="FFFFC000")


def style_header_row(ws, row, col_count):
    """將指定列設為表頭樣式：深色底、白色字體。"""
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
